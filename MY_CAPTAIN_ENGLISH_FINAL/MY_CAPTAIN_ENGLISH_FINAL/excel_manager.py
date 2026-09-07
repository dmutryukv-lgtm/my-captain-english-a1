from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

REQUIRED_SHEETS = [
    "README", "USERS", "LEVELS", "LESSONS", "LESSON_STEPS",
    "WORDS", "PRONUNCIATION", "PHRASES", "GRAMMAR", "LISTENING",
    "SPEAKING_SCENARIOS", "USER_WORDS", "USER_MISTAKES", "PROGRESS",
    "DAILY_GOALS", "ACHIEVEMENTS", "USER_ACHIEVEMENTS", "SETTINGS"
]

class ExcelDB:
    def __init__(self, path):
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(f"Excel не знайдено: {self.path}")
        self._check()

    def _check(self):
        wb = load_workbook(self.path, read_only=True, data_only=True)
        missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
        wb.close()
        if missing:
            raise ValueError("У Excel відсутні листи: " + ", ".join(missing))

    def read(self, sheet):
        wb = load_workbook(self.path, read_only=True, data_only=True)
        if sheet not in wb.sheetnames:
            wb.close()
            return []
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(x).strip() if x is not None else "" for x in rows[0]]
        result = []
        for row in rows[1:]:
            item = {}
            empty = True
            for i, h in enumerate(headers):
                value = row[i] if i < len(row) else None
                if value not in (None, ""):
                    empty = False
                item[h] = value
            if not empty:
                result.append(item)
        return result

    def append(self, sheet, values):
        wb = load_workbook(self.path)
        if sheet not in wb.sheetnames:
            wb.close()
            raise ValueError(f"Лист {sheet} не знайдено")
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        ws.append([values.get(h) for h in headers])
        wb.save(self.path)
        wb.close()

    def update_matching(self, sheet, match, updates):
        wb = load_workbook(self.path)
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        positions = {h: i + 1 for i, h in enumerate(headers)}
        changed = False
        for r in range(2, ws.max_row + 1):
            ok = all(ws.cell(r, positions[k]).value == v for k, v in match.items())
            if ok:
                for k, v in updates.items():
                    if k in positions:
                        ws.cell(r, positions[k]).value = v
                changed = True
                break
        if changed:
            wb.save(self.path)
        wb.close()
        return changed

    def upsert_by_keys(self, sheet, keys, values):
        wb = load_workbook(self.path)
        ws = wb[sheet]
        headers = [c.value for c in ws[1]]
        positions = {h: i + 1 for i, h in enumerate(headers)}
        found = None
        for r in range(2, ws.max_row + 1):
            if all(ws.cell(r, positions[k]).value == v for k, v in keys.items()):
                found = r
                break

        if found is None:
            ws.append([values.get(h) for h in headers])
        else:
            for k, v in values.items():
                if k in positions:
                    ws.cell(found, positions[k]).value = v

        wb.save(self.path)
        wb.close()
